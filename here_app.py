# import needed modules. tkinter[GUI interface]
# and python-docx[Interact with the document]
# will be the modules used
# here_app.py
import os  # Import the os module
from docx import Document
from gui_module import create_app

# Path of file. Change later [refer to notes on Notion]
doc_path = 'E://NO//CODE//projects//Here-App//class_attendance.docx'
doc = Document(doc_path)
table = doc.tables[0]
student_row_number = 1

def action_on_input(value, label):
    global student_row_number
    if student_row_number < len(table.rows):
        student_row_number += 1  # increment used for next student
        update_student()  # update label with next student name
    else:
        label.config(text="Done! You may now close the application.")

# Get the document name from the document path
doc_name = os.path.basename(doc_path)

# Pass the document name as the second argument to create_app
app = create_app(action_on_input, doc_name)

def update_student():
    global student_row_number
    if student_row_number < len(table.rows):
        student_name = table.cell(student_row_number, 0).text
        app.update_label(student_name)
    else:
        app.update_label("Done! You may now close the application.")

update_student()
app.run()
